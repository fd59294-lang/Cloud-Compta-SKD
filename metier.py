"""Règles métier sans interface : opérations, soldes, synthèses, banque, budget, imports, sauvegarde Excel.

Toutes les fonctions reçoivent et renvoient des DataFrame : elles se testent sans Streamlit ni Google.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
import unicodedata
import uuid
from datetime import date, datetime, timedelta

import pandas as pd

from referentiel import (CATEGORIES_DEFAUT, COLONNES_IMPORT, CORRESPONDANCE_CATEGORIES, MODES_PAIEMENT,
                         MOIS_COURTS, SENS, TABLES)


# ------------------------------------------------------------------ formats
def eur(x: float, signe: bool = False) -> str:
    x = float(x or 0)
    s = f"{abs(x):,.2f}".replace(",", " ").replace(".", ",") + " €"
    if round(x, 2) < 0:
        return "−" + s
    return ("+" + s) if signe and x > 0 else s


def nombres(serie: pd.Series) -> pd.Series:
    """Convertit une colonne (texte, nombres, vide) en float64, quel que soit le type d'origine."""
    valeurs = ["" if v is None else str(v).replace(" ", "").replace("\xa0", "").replace(" ", "")
               .replace("€", "").replace(",", ".").strip() for v in serie.tolist()]
    return pd.Series(pd.to_numeric(pd.Series(valeurs, index=serie.index, dtype=object), errors="coerce"),
                     index=serie.index, dtype="float64")


def nombre(v, defaut: float | None = None) -> float | None:
    s = pd.Series([v])
    n = nombres(s).iloc[0]
    return defaut if pd.isna(n) else float(n)


def lire_date(v) -> date | None:
    if v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2}|\d{4})$", s)
    if m:
        a = int(m.group(3))
        try:
            return date(a + 2000 if a < 100 else a, int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    try:
        d = pd.to_datetime(s, errors="coerce")
        return None if pd.isna(d) else d.date()
    except Exception:
        return None


def oui(v) -> bool:
    return str(v).strip().lower() in ("oui", "true", "vrai", "1", "x", "r", "yes")


def sans_accent(s) -> str:
    return unicodedata.normalize("NFD", str(s or "").lower()).encode("ascii", "ignore").decode()


def nouvel_id() -> str:
    return uuid.uuid4().hex[:12]


def maintenant() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ------------------------------------------------------------------ exercices
def exercice_de(d: date, mois_debut: int = 9) -> int:
    """Exercice = année de début. Avec mois_debut=9 : sept. 2025 → août 2026 = exercice 2025."""
    return d.year if d.month >= mois_debut else d.year - 1


def libelle_exercice(ex: int, mois_debut: int = 9) -> str:
    return str(ex) if mois_debut == 1 else f"{ex}-{ex + 1}"


def bornes_exercice(ex: int, mois_debut: int = 9) -> tuple[date, date]:
    debut = date(int(ex), mois_debut, 1)
    fin = date(int(ex) + 1, mois_debut, 1) - timedelta(days=1)
    return debut, fin


def libelle_categorie(code, nom) -> str:
    return f"{code} — {nom}"


def code_categorie(libelle: str) -> str:
    return str(libelle).split(" — ")[0].strip()


def nom_categorie(libelle: str) -> str:
    s = str(libelle)
    return s.split(" — ", 1)[1] if " — " in s else s


# ------------------------------------------------------------------ opérations
def typer_operations(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    if df.empty:
        df = pd.DataFrame(columns=TABLES["operations"])
    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    df["montant"] = nombres(df["montant"]).fillna(0.0)
    df["exercice"] = nombres(df["exercice"]).fillna(0).astype(int)
    df["signe"] = df["montant"].where(df["sens"] == "Recette", -df["montant"])
    df["rapproche_b"] = df["rapproche"].map(oui)
    if "releve" not in df:
        df["releve"] = ""
    # Date du relevé sur lequel l'écriture est pointée (à défaut : sa propre date)
    df["releve_d"] = pd.to_datetime(df["releve"].where(df["releve"].astype(str).str.strip() != "", None),
                                    errors="coerce").fillna(df["date"])
    return df


def nouvelle_operation(d: date, sens: str, categorie: str, libelle: str, montant: float, source: str,
                       mode: str, tiers: str, piece: str, notes: str, utilisateur: str, mois_debut: int = 9,
                       rapproche: bool = False, origine: str = "", justificatif: str = "",
                       releve: date | None = None) -> dict:
    if sens not in SENS:
        raise ValueError("Sens inconnu")
    return {
        "id": nouvel_id(), "date": d.isoformat(), "exercice": exercice_de(d, mois_debut), "sens": sens,
        "categorie": categorie, "libelle": str(libelle).strip(), "montant": round(abs(float(montant)), 2),
        "source": "Espèces" if source == "Espèces" else "Banque", "mode_paiement": mode or "",
        "tiers": str(tiers or "").strip(), "piece": str(piece or "").strip(), "notes": str(notes or "").strip(),
        "rapproche": "oui" if rapproche else "non", "justificatif": justificatif, "origine": origine,
        "saisi_par": utilisateur, "saisi_le": maintenant(),
        "releve": releve.isoformat() if (rapproche and releve) else "",
    }


def soldes(ops: pd.DataFrame, jusqu_au: date | None = None, avant: date | None = None) -> dict:
    """Soldes banque et caisse (toutes écritures jusqu'à une date incluse, ou strictement avant une date)."""
    x, p = ops, ops
    if jusqu_au is not None:
        x = x[x["date"] <= pd.Timestamp(jusqu_au)]
        p = p[p["releve_d"] <= pd.Timestamp(jusqu_au)]
    if avant is not None:
        x = x[x["date"] < pd.Timestamp(avant)]
        p = p[p["releve_d"] < pd.Timestamp(avant)]
    banque = x.loc[x["source"] != "Espèces", "signe"].sum()
    caisse = x.loc[x["source"] == "Espèces", "signe"].sum()
    pointe = p.loc[(p["source"] != "Espèces") & p["rapproche_b"], "signe"].sum()
    return {"banque": round(float(banque), 2), "caisse": round(float(caisse), 2), "pointe": round(float(pointe), 2)}


def totaux(ops: pd.DataFrame, exercice: int) -> dict:
    x = ops[ops["exercice"] == int(exercice)]
    r = x.loc[x["sens"] == "Recette", "montant"].sum()
    d = x.loc[x["sens"] == "Dépense", "montant"].sum()
    return {"recettes": round(float(r), 2), "depenses": round(float(d), 2), "resultat": round(float(r - d), 2),
            "nb": int(len(x))}


def synthese_categories(ops: pd.DataFrame, exercice: int) -> pd.DataFrame:
    x = ops[ops["exercice"] == int(exercice)]
    if x.empty:
        return pd.DataFrame(columns=["sens", "categorie", "montant"])
    return (x.groupby(["sens", "categorie"], as_index=False)["montant"].sum()
             .sort_values(["sens", "montant"], ascending=[False, False]))


def synthese_mensuelle(ops: pd.DataFrame, exercice: int, mois_debut: int = 9) -> pd.DataFrame:
    """Tableau « Synthèse des opérations » de l'Excel : catégories en lignes, mois en colonnes (montants signés)."""
    x = ops[ops["exercice"] == int(exercice)].copy()
    ordre = [MOIS_COURTS[(mois_debut - 1 + k) % 12] for k in range(12)]
    if x.empty:
        return pd.DataFrame(columns=["Sens", "Catégorie"] + ordre + ["Total"])
    x["mois"] = x["date"].dt.month.map(lambda m: MOIS_COURTS[int(m) - 1])
    t = x.pivot_table(index=["sens", "categorie"], columns="mois", values="signe", aggfunc="sum", fill_value=0)
    t = t.reindex(columns=ordre, fill_value=0)
    t["Total"] = t.sum(axis=1)
    t = t.reset_index().rename(columns={"sens": "Sens", "categorie": "Catégorie"})
    t = t.sort_values(["Sens", "Catégorie"], ascending=[False, True])
    total = {"Sens": "", "Catégorie": "Total général", **{m: t[m].sum() for m in ordre + ["Total"]}}
    return pd.concat([t, pd.DataFrame([total])], ignore_index=True)


def mensuel(ops: pd.DataFrame, exercice: int, mois_debut: int = 9) -> pd.DataFrame:
    """Recettes et dépenses par mois de l'exercice (pour les graphiques)."""
    x = ops[ops["exercice"] == int(exercice)]
    lignes = []
    for k in range(12):
        m = (mois_debut - 1 + k) % 12 + 1
        y = x[x["date"].dt.month == m]
        lignes.append({"ordre": k, "mois": MOIS_COURTS[m - 1],
                       "Recettes": float(y.loc[y["sens"] == "Recette", "montant"].sum()),
                       "Dépenses": float(y.loc[y["sens"] == "Dépense", "montant"].sum())})
    return pd.DataFrame(lignes)


# ------------------------------------------------------------------ catégories
def table_categories(df_categories: pd.DataFrame | None = None) -> pd.DataFrame:
    if df_categories is None or df_categories.empty:
        return pd.DataFrame(CATEGORIES_DEFAUT, columns=["code", "nom", "sens"]).assign(actif="oui")
    return df_categories


def trouver_categorie(texte, sens: str, df_categories: pd.DataFrame | None = None) -> tuple[str, bool]:
    """Libellé « code — nom » pour un texte libre (code, nom, ancienne catégorie). (libellé, reconnu ?)"""
    cats = table_categories(df_categories)
    t = sans_accent(texte).strip()
    if t:
        for r in cats.itertuples():
            if t in (sans_accent(r.code), sans_accent(r.nom), sans_accent(libelle_categorie(r.code, r.nom))):
                return libelle_categorie(r.code, r.nom), True
        for ancien, code in CORRESPONDANCE_CATEGORIES.items():
            if t == sans_accent(ancien):
                ligne = cats[cats["code"].astype(str) == code]
                if not ligne.empty:
                    return libelle_categorie(code, ligne.iloc[0]["nom"]), True
    code = "R99" if sens == "Recette" else "D99"
    ligne = cats[cats["code"].astype(str) == code]
    nom = ligne.iloc[0]["nom"] if not ligne.empty else f"À classer ({sens.lower()})"
    return libelle_categorie(code, nom), False


# ------------------------------------------------------------------ doublons
def empreinte(d, montant, sens, libelle) -> str:
    cle = f"{pd.Timestamp(d).date()}|{sens}|{round(float(montant), 2)}|{sans_accent(libelle).strip()}"
    return hashlib.sha1(cle.encode()).hexdigest()


def empreintes_existantes(ops: pd.DataFrame) -> set:
    return {empreinte(r.date, r.montant, r.sens, r.libelle) for r in ops.itertuples() if pd.notna(r.date)}


# ------------------------------------------------------------------ import de l'Excel d'origine (.xlsm)
def lire_livre_excel(contenu: bytes) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    """Lit la feuille « Livre de comptes » (et « Opérations en attente ») du classeur Excel d'origine.

    Colonnes de l'Excel : R (rapproché) | Dates | Montants (signés) | E (espèces) | Catégories | Tiers | Notes.
    Renvoie (écritures, opérations en attente, messages)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contenu), data_only=True, read_only=True)
    messages = []

    def lire(feuille, decal):
        lignes, entete = [], None
        ws = wb[feuille]
        for i, r in enumerate(ws.iter_rows(max_col=decal + 7, values_only=True)):
            r = list(r) + [None] * 8
            if entete is None:
                if any(str(v).strip().lower() == "dates" for v in r if v is not None):
                    entete = i
                continue
            d, m = r[decal + 1], r[decal + 2]
            if not isinstance(d, (datetime, date)) or not isinstance(m, (int, float)) or m == 0:
                continue
            lignes.append({"rapproche": r[decal] not in (None, "", 0), "date": lire_date(d), "montant": float(m),
                           "especes": str(r[decal + 3] or "").strip().upper() == "E",
                           "categorie": str(r[decal + 4] or "").strip(), "tiers": str(r[decal + 5] or "").strip(),
                           "notes": str(r[decal + 6] or "").strip()})
        return pd.DataFrame(lignes)

    if "Livre de comptes" not in wb.sheetnames:
        return pd.DataFrame(), pd.DataFrame(), ["Feuille « Livre de comptes » introuvable dans ce fichier."]
    livre = lire("Livre de comptes", 0)
    attente = pd.DataFrame()
    if "Opérations en attente" in wb.sheetnames:
        attente = _lire_attente(wb["Opérations en attente"])
    messages.append(f"{len(livre)} écriture(s) lue(s) dans le livre de comptes, {len(attente)} opération(s) en attente.")
    return livre, attente, messages


def _lire_attente(ws) -> pd.DataFrame:
    """Opérations en attente : Dates | Montants | E | Catégories | Tiers | Notes (sans colonne R)."""
    lignes, entete = [], None
    for i, r in enumerate(ws.iter_rows(max_col=8, values_only=True)):
        r = list(r) + [None] * 8
        if entete is None:
            if any(str(v).strip().lower() == "dates" for v in r if v is not None):
                entete = i
                debut = [str(v).strip().lower() if v else "" for v in r].index("dates")
            continue
        d, m = r[debut], r[debut + 1]
        if not isinstance(d, (datetime, date)) or not isinstance(m, (int, float)) or m == 0:
            continue
        lignes.append({"date": lire_date(d), "montant": float(m),
                       "especes": str(r[debut + 2] or "").strip().upper() == "E",
                       "categorie": str(r[debut + 3] or "").strip(), "tiers": str(r[debut + 4] or "").strip(),
                       "notes": str(r[debut + 5] or "").strip()})
    return pd.DataFrame(lignes)


def convertir_livre(livre: pd.DataFrame, df_categories: pd.DataFrame, utilisateur: str, mois_debut: int,
                    existantes: set) -> pd.DataFrame:
    """Transforme les lignes de l'Excel en écritures du programme, avec statut (OK / Doublon / À classer)."""
    lignes = []
    for r in livre.to_dict("records"):
        sens = "Recette" if r["montant"] > 0 else "Dépense"
        cat, reconnu = trouver_categorie(r["categorie"], sens, df_categories)
        libelle = r["notes"] or r["categorie"] or r["tiers"] or "Opération"
        notes = "" if reconnu else f"Catégorie d'origine : {r['categorie']}"
        op = nouvelle_operation(r["date"], sens, cat, libelle, abs(r["montant"]),
                                "Espèces" if r["especes"] else "Banque", "", r["tiers"], "", notes,
                                f"reprise ({utilisateur})", mois_debut, rapproche=bool(r.get("rapproche")),
                                origine="reprise Excel")
        op["statut"] = ("Doublon" if empreinte(r["date"], abs(r["montant"]), sens, libelle) in existantes
                        else ("OK" if reconnu else "À classer"))
        lignes.append(op)
    return pd.DataFrame(lignes)


# ------------------------------------------------------------------ import de la version Apps Script (export .xlsx)
def lire_export_apps_script(contenu: bytes) -> dict[str, pd.DataFrame]:
    """Lit l'export Excel du Google Sheets « Comptabilité_Association » (version Apps Script)."""
    feuilles = pd.read_excel(io.BytesIO(contenu), sheet_name=None, dtype=object)
    return {k: v for k, v in feuilles.items()}


def convertir_operations_apps_script(df: pd.DataFrame, df_categories: pd.DataFrame, utilisateur: str,
                                     mois_debut: int, existantes: set) -> pd.DataFrame:
    """Onglet Operations_Historique : Date | Libelle | Montant (signé) | Type | Categorie | Tiers | Notes |
    Numero_Piece | Source | Rapproche | Solde_Cumule | Justificatif."""
    col = {sans_accent(c).replace(" ", "_"): c for c in df.columns}

    def v(r, nom, defaut=""):
        c = col.get(nom)
        x = r.get(c) if c else None
        return defaut if x is None or (isinstance(x, float) and pd.isna(x)) else x

    lignes = []
    for r in df.to_dict("records"):
        d, m = lire_date(v(r, "date")), nombre(v(r, "montant"))
        if not d or not m:
            continue
        type_ = sans_accent(v(r, "type"))
        sens = "Recette" if (type_.startswith("recette") or (not type_ and m > 0) or (m > 0 and not
                             type_.startswith("depense"))) else "Dépense"
        cat, reconnu = trouver_categorie(v(r, "categorie"), sens, df_categories)
        libelle = str(v(r, "libelle") or v(r, "categorie") or "Opération")
        source = "Espèces" if sans_accent(v(r, "source")).startswith("espece") else "Banque"
        notes = str(v(r, "notes"))
        if not reconnu and v(r, "categorie"):
            notes = (notes + " · " if notes else "") + f"Catégorie d'origine : {v(r, 'categorie')}"
        op = nouvelle_operation(d, sens, cat, libelle, abs(m), source, "", str(v(r, "tiers")),
                                str(v(r, "numero_piece")), notes, f"reprise ({utilisateur})", mois_debut,
                                rapproche=oui(v(r, "rapproche")), origine="reprise Apps Script",
                                justificatif=str(v(r, "justificatif")))
        op["statut"] = ("Doublon" if empreinte(d, abs(m), sens, libelle) in existantes
                        else ("OK" if reconnu else "À classer"))
        lignes.append(op)
    return pd.DataFrame(lignes)


# ------------------------------------------------------------------ import générique (modèle Excel)
def modele_import_excel() -> bytes:
    exemple = pd.DataFrame([
        ["15/09/2026", "Recette", "756", "Cotisation DUPONT Léo", 129.00, "Banque", "Chèque", "DUPONT", "REC-2026-001", "",
         "non", ""],
        ["09/10/2026", "Dépense", "Frais bancaires", "Frais de tenue de compte", 3.88, "Banque", "Prélèvement",
         "Crédit Mutuel", "", "", "oui", "31/10/2026"],
        ["03/02/2027", "Recette", "Subventions", "Subvention mairie", 500, "Banque", "Virement", "Mairie de Haussy",
         "", "", "non", ""],
    ], columns=COLONNES_IMPORT)
    aide = pd.DataFrame({
        "Colonne": COLONNES_IMPORT,
        "Obligatoire": ["Oui", "Oui*", "Non", "Oui", "Oui", "Non", "Non", "Non", "Non", "Non", "Non", "Non"],
        "Ce qu'il faut mettre": [
            "Date de l'opération (JJ/MM/AAAA)",
            "Recette ou Dépense (*facultatif si le montant est négatif pour une dépense)",
            "Code (ex. 627) ou nom de catégorie ; vide = « À classer »",
            "Description courte de l'opération",
            "Montant en euros (virgule ou point acceptés)",
            "Banque ou Espèces (vide = Banque)",
            "Chèque, Virement, Espèces, Pass'Sport, Carte bancaire, Prélèvement, Autre",
            "Fournisseur, adhérent, donateur…",
            "Numéro de facture, de ticket ou de reçu",
            "Informations complémentaires",
            "oui si l'opération figure déjà sur un relevé bancaire",
            "Date de fin du relevé sur lequel elle figure (facultatif)",
        ],
    })
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        exemple.to_excel(w, sheet_name="Opérations", index=False)
        aide.to_excel(w, sheet_name="Mode d'emploi", index=False)
        mettre_en_forme(w)
    return buf.getvalue()


def convertir_modele(contenu: bytes, df_categories: pd.DataFrame, utilisateur: str, mois_debut: int,
                     existantes: set) -> tuple[pd.DataFrame, list[str]]:
    try:
        df = pd.read_excel(io.BytesIO(contenu), sheet_name=0, dtype=object)
    except Exception as e:
        return pd.DataFrame(), [f"Fichier illisible : {e}"]
    noms = {sans_accent(c).strip(): c for c in df.columns}
    manquantes = [c for c in ("Date", "Libellé", "Montant") if sans_accent(c) not in noms]
    if manquantes:
        return pd.DataFrame(), ["Colonnes obligatoires absentes : " + ", ".join(manquantes)
                                + ". Utilisez le modèle d'import."]

    def v(r, c):
        x = r.get(noms.get(sans_accent(c)))
        return "" if x is None or (isinstance(x, float) and pd.isna(x)) else x

    lignes, erreurs = [], []
    for i, r in enumerate(df.to_dict("records"), start=2):
        d, m = lire_date(v(r, "Date")), nombre(v(r, "Montant"))
        if not d or m is None or m == 0:
            if any(str(x).strip() for x in r.values() if x is not None and not (isinstance(x, float) and pd.isna(x))):
                erreurs.append(f"Ligne {i} ignorée : date ou montant illisible.")
            continue
        s = sans_accent(v(r, "Sens"))
        sens = "Recette" if s.startswith("rec") else "Dépense" if s.startswith("dep") else (
            "Recette" if m > 0 else "Dépense")
        cat, reconnu = trouver_categorie(v(r, "Catégorie"), sens, df_categories)
        libelle = str(v(r, "Libellé")).strip() or "Opération"
        mode = str(v(r, "Mode de paiement")).strip()
        op = nouvelle_operation(d, sens, cat, libelle, abs(m),
                                "Espèces" if sans_accent(v(r, "Source")).startswith("espece") else "Banque",
                                mode if mode in MODES_PAIEMENT else "", str(v(r, "Tiers")), str(v(r, "N° pièce")),
                                str(v(r, "Notes")), f"import ({utilisateur})", mois_debut,
                                rapproche=oui(v(r, "Rapproché")), origine="import Excel",
                                releve=lire_date(v(r, "Relevé du")))
        op["statut"] = ("Doublon" if empreinte(d, abs(m), sens, libelle) in existantes
                        else ("OK" if reconnu else "À classer"))
        lignes.append(op)
    return pd.DataFrame(lignes), erreurs


# ------------------------------------------------------------------ relevé bancaire CSV
def lire_releve(contenu: bytes, nom_fichier: str = "") -> pd.DataFrame:
    """Lit un relevé exporté par la banque : CSV, ou Excel (.xlsx, export du Crédit Mutuel par exemple).

    Le solde de fin éventuel (« Solde au 31/10/2026 : 213,46 ») est rendu dans .attrs["solde_fin"] = (date, solde)."""
    if nom_fichier.lower().endswith((".xlsx", ".xls")) or contenu[:2] == b"PK":
        return lire_releve_excel(contenu)
    return lire_releve_csv(contenu)


def lire_releve_csv(contenu: bytes) -> pd.DataFrame:
    """Lit un relevé CSV de banque : en-tête avec Date + Libellé + (Débit/Crédit ou Montant)."""
    try:
        texte = contenu.decode("utf-8-sig")
    except UnicodeDecodeError:
        texte = contenu.decode("latin-1")
    debut = "\n".join(texte.splitlines()[:15])
    sep = ";" if debut.count(";") >= debut.count(",") else ("\t" if "\t" in debut else ",")
    return _lignes_releve(list(csv.reader(io.StringIO(texte), delimiter=sep)))


def lire_releve_excel(contenu: bytes) -> pd.DataFrame:
    """Lit un relevé Excel : la première feuille qui contient un en-tête Date + Libellé + Débit/Crédit."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contenu), data_only=True, read_only=True)
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        rows = [["" if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]
        try:
            return _lignes_releve(rows)
        except ValueError:
            continue
    raise ValueError("Aucune feuille du fichier ne contient d'en-tête Date, Libellé, Débit/Crédit.")


def _lignes_releve(rows: list[list]) -> pd.DataFrame:
    h = None
    for k, r in enumerate(rows[:40]):
        c = [sans_accent(x).strip() for x in r]
        i_date = next((j for j, x in enumerate(c) if x.startswith("date") and "valeur" not in x), None)
        i_deb = next((j for j, x in enumerate(c) if "debit" in x), None)
        i_cre = next((j for j, x in enumerate(c) if "credit" in x), None)
        i_mt = next((j for j, x in enumerate(c) if "montant" in x), None)
        if i_date is not None and (i_mt is not None or i_deb is not None or i_cre is not None):
            i_lib = next((j for j, x in enumerate(c) if re.search(r"libelle|operation|intitule|description|detail", x)),
                         None)
            h = k
            break
    if h is None:
        raise ValueError("Format non reconnu : il faut une ligne d'en-tête avec Date, Libellé et Débit/Crédit "
                         "(ou Montant). Exportez le relevé en CSV ou en Excel depuis votre banque en ligne.")
    lignes, solde_fin = [], None
    for r in rows[h + 1:]:
        r = list(r) + [""] * 10
        texte_ligne = " ".join(str(x) for x in r if x not in ("", None))
        m_solde = re.search(r"solde au (\d{1,2}/\d{1,2}/\d{4})", sans_accent(texte_ligne))
        if m_solde:
            valeurs = [x for x in r if isinstance(x, (int, float)) and not isinstance(x, bool)]
            if valeurs:
                solde_fin = (lire_date(m_solde.group(1)), round(float(valeurs[-1]), 2))
            continue
        d = lire_date(r[i_date])
        if not d:
            continue
        m = None
        if i_mt is not None and str(r[i_mt]).strip():
            m = nombre(r[i_mt])
        else:
            cr = nombre(r[i_cre]) if i_cre is not None else None
            db = nombre(r[i_deb]) if i_deb is not None else None
            if cr:
                m = abs(cr)
            elif db:
                m = -abs(db)
        if not m:
            continue
        lignes.append({"date": d, "libelle": re.sub(r"\s+", " ", str(r[i_lib])).strip() if i_lib is not None else "",
                       "montant": round(m, 2)})
    df = pd.DataFrame(lignes, columns=["date", "libelle", "montant"])
    df.attrs["solde_fin"] = solde_fin
    return df


_MOTS_VIDES = {"prlv", "prelevement", "vir", "virement", "sepa", "carte", "chq", "cheque", "remise", "des", "les",
               "pour", "par", "ref", "inst", "recu", "emis", "facture", "frais", "paiement", "achat", "date",
               "association", "the"}


def mots_cles(s) -> list[str]:
    return [w for w in re.sub(r"[^a-z ]", " ", sans_accent(s)).split() if len(w) > 2 and w not in _MOTS_VIDES]


def analyser_releve(releve: pd.DataFrame, ops: pd.DataFrame, categories_dispo: list[str]) -> pd.DataFrame:
    """Pour chaque ligne du relevé : « À rapprocher » (déjà saisie), « Déjà rapprochée » ou « Nouvelle »,
    avec une catégorie proposée d'après les habitudes de saisie."""
    habitudes = {}
    for r in ops.sort_values("date").itertuples():
        for w in mots_cles(f"{r.libelle} {r.tiers}"):
            habitudes[w] = r.categorie
    banque = ops[ops["source"] != "Espèces"]
    pris = set()
    lignes = []
    for b in releve.to_dict("records"):
        meilleur, ecart_min, deja, ecart_deja = None, 999, None, 999
        cible = banque[(banque["signe"].round(2) == round(b["montant"], 2))]
        for r in cible.itertuples():
            if r.id in pris or pd.isna(r.date):
                continue
            ecart = abs((r.date.date() - b["date"]).days)
            if ecart > 10:
                continue
            if r.rapproche_b:
                if ecart < ecart_deja:
                    deja, ecart_deja = r, ecart
            elif ecart < ecart_min:
                meilleur, ecart_min = r, ecart
        trouve = meilleur or deja
        if trouve is not None:
            pris.add(trouve.id)
        sens = "Recette" if b["montant"] > 0 else "Dépense"
        cat = categorie_par_regle(b["libelle"], sens, categories_dispo) or \
            next((habitudes[w] for w in mots_cles(b["libelle"]) if w in habitudes), "")
        if not cat or cat not in categories_dispo or (cat in categories_dispo and _sens_categorie(cat) != sens):
            cat = next((c for c in categories_dispo if c.startswith("627 ")), "") if (
                sens == "Dépense" and re.search(r"frais|commission|cotis.*carte|tenue|abonnement", sans_accent(b["libelle"]))
            ) else ""
        lignes.append({"importer": deja is None or meilleur is not None, "date": b["date"], "libelle": b["libelle"],
                       "montant": b["montant"],
                       "resultat": "À rapprocher" if meilleur is not None else ("Déjà rapprochée" if deja is not None
                                                                                 else "Nouvelle"),
                       "operation_id": trouve.id if trouve is not None else "",
                       "ecriture": f"{trouve.date:%d/%m/%Y} · {trouve.libelle}" if trouve is not None else "",
                       "categorie": cat})
    return pd.DataFrame(lignes)


# Libellés habituels des relevés (Crédit Mutuel) -> code de catégorie, selon le sens
REGLES_BANQUE = [
    (r"soutien asso|retrocession|remb.*frais", "Recette", "7581"),
    (r"fact sgt|frais|commission|cotis.*carte|tenue de compte|abonnement|agios", "Dépense", "627"),
    (r"federation francaise|ffk|ffkda|licence", "Dépense", "6281"),
    (r"ligue|comite departemental|cdk|affiliation", "Dépense", "6282"),
    (r"maif|macif|assurance|allianz|axa", "Dépense", "616"),
    (r"rem chq|remise ch|licence|cotis|adhesion|inscription|pass.?sport|pass psp|asp agence|drfip", "Recette", "756"),
    (r"subvention|mairie|commune|departement|region|ans ", "Recette", "74"),
    (r"don ", "Recette", "754"),
]


def categorie_par_regle(libelle: str, sens: str, categories_dispo: list[str]) -> str:
    t = sans_accent(libelle) + " "
    for motif, s, code in REGLES_BANQUE:
        if s == sens and re.search(motif, t):
            cat = next((c for c in categories_dispo if code_categorie(c) == code), "")
            if cat:
                return cat
    return ""


def _sens_categorie(libelle: str) -> str:
    code = code_categorie(libelle)
    for c, _, s in CATEGORIES_DEFAUT:
        if c == code:
            return s
    return "Recette" if code.startswith("7") or code.startswith("R") or code.startswith("1") else "Dépense"


# ------------------------------------------------------------------ relevés validés (verrouillage)
def dernier_releve_valide(releves: pd.DataFrame) -> tuple[date | None, float]:
    """Date et solde du dernier relevé validé (écart nul). (None, 0) s'il n'y en a pas."""
    if releves is None or releves.empty or "valide" not in releves:
        return None, 0.0
    v = releves[releves["valide"].map(oui)]
    if v.empty:
        return None, 0.0
    v = v.assign(d=v["date_releve"].map(lire_date)).dropna(subset=["d"]).sort_values("d")
    if v.empty:
        return None, 0.0
    return v.iloc[-1]["d"], float(nombre(v.iloc[-1]["solde_releve"], 0.0))


def verrouillees(ops: pd.DataFrame, date_verrou: date | None) -> pd.Series:
    """Écritures bancaires pointées sur un relevé déjà validé : elles ne se modifient plus."""
    if date_verrou is None or ops.empty:
        return pd.Series(False, index=ops.index)
    return (ops["source"] != "Espèces") & ops["rapproche_b"] & (ops["releve_d"] <= pd.Timestamp(date_verrou))


# ------------------------------------------------------------------ opérations en attente
def serie_mensuelle(premiere: date, nombre_mois: int) -> list[date]:
    dates = []
    for k in range(nombre_mois):
        a, m = premiere.year + (premiere.month - 1 + k) // 12, (premiere.month - 1 + k) % 12 + 1
        jour = min(premiere.day, [31, 29 if a % 4 == 0 else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1])
        dates.append(date(a, m, jour))
    return dates


# ------------------------------------------------------------------ budget
def budget_compare(ops: pd.DataFrame, budget: pd.DataFrame, exercice: int) -> pd.DataFrame:
    """Prévu / réalisé / réalisé N-1 par catégorie."""
    b = budget[nombres(budget["exercice"]) == int(exercice)] if not budget.empty else budget
    lignes = {}
    for r in (b.to_dict("records") if not b.empty else []):
        lignes[(r["sens"], r["categorie"])] = {"sens": r["sens"], "categorie": r["categorie"],
                                              "prevu": nombre(r["prevu"], 0.0), "commentaire": r.get("commentaire", "")}
    for ex, col in ((int(exercice), "realise"), (int(exercice) - 1, "realise_n1")):
        for r in synthese_categories(ops, ex).to_dict("records"):
            k = (r["sens"], r["categorie"])
            lignes.setdefault(k, {"sens": r["sens"], "categorie": r["categorie"], "prevu": 0.0, "commentaire": ""})
            lignes[k][col] = float(r["montant"])
    df = pd.DataFrame(list(lignes.values()), columns=["sens", "categorie", "prevu", "realise", "realise_n1",
                                                      "commentaire"])
    df[["prevu", "realise", "realise_n1"]] = df[["prevu", "realise", "realise_n1"]].fillna(0.0)
    df["ecart"] = df["realise"] - df["prevu"]
    return df.sort_values(["sens", "categorie"], ascending=[False, True]).reset_index(drop=True)


def proposer_budget(ops: pd.DataFrame, exercice: int, hausse_pct: float = 0.0) -> pd.DataFrame:
    """Budget proposé = réalisé de l'exercice précédent (+ hausse), arrondi à la dizaine supérieure."""
    s = synthese_categories(ops, int(exercice) - 1)
    if s.empty:
        return pd.DataFrame(columns=["sens", "categorie", "prevu"])
    s["prevu"] = (s["montant"] * (1 + hausse_pct / 100) / 10).apply(lambda x: max(0, -(-x // 1)) * 10)
    return s[["sens", "categorie", "prevu"]]


# ------------------------------------------------------------------ montant en lettres
_U = ["zéro", "un", "deux", "trois", "quatre", "cinq", "six", "sept", "huit", "neuf", "dix", "onze", "douze",
      "treize", "quatorze", "quinze", "seize", "dix-sept", "dix-huit", "dix-neuf"]
_D = ["", "", "vingt", "trente", "quarante", "cinquante", "soixante", "soixante", "quatre-vingt", "quatre-vingt"]


def _moins_de_cent(x: int) -> str:
    if x < 20:
        return _U[x]
    dz, un = divmod(x, 10)
    if dz in (7, 9):
        return _D[dz] + (" et " if dz == 7 and un == 1 else "-") + _U[10 + un]
    if un == 0:
        return _D[dz] + ("s" if dz == 8 else "")
    return _D[dz] + (" et " if un == 1 and dz != 8 else "-") + _U[un]


def _moins_de_mille(x: int) -> str:
    c, r = divmod(x, 100)
    s = ""
    if c:
        s = (_U[c] + " cent" if c > 1 else "cent") + ("s" if c > 1 and r == 0 else "")
    if r:
        s += (" " if s else "") + _moins_de_cent(r)
    return s


def _entier(x: int) -> str:
    if x == 0:
        return "zéro"
    mil, reste = divmod(x, 1_000_000)
    mille, r = divmod(reste, 1000)
    s = ""
    if mil:
        s += _moins_de_mille(mil) + (" millions" if mil > 1 else " million")
    if mille:
        s += (" " if s else "") + ("mille" if mille == 1 else _moins_de_mille(mille).removesuffix("s") + " mille")
    if r:
        s += (" " if s else "") + _moins_de_mille(r)
    return s


def en_lettres(montant: float) -> str:
    total = round(abs(float(montant)) * 100)
    e, c = divmod(total, 100)
    s = _entier(e) + (" euros" if e > 1 else " euro")
    if c:
        s += " et " + _entier(c) + (" centimes" if c > 1 else " centime")
    return s


# ------------------------------------------------------------------ sauvegarde Excel
def export_excel(tables: dict[str, pd.DataFrame], mois_debut: int = 9) -> bytes:
    ops = typer_operations(tables["operations"])
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        vue = ops.sort_values("date").drop(columns=["signe", "rapproche_b", "releve_d"])
        vue["date"] = vue["date"].dt.date
        vue.rename(columns={"id": "Identifiant", "date": "Date", "exercice": "Exercice", "sens": "Sens",
                            "categorie": "Catégorie", "libelle": "Libellé", "montant": "Montant", "source": "Source",
                            "mode_paiement": "Mode de paiement", "tiers": "Tiers", "piece": "N° pièce",
                            "notes": "Notes", "rapproche": "Rapproché", "justificatif": "Justificatif",
                            "origine": "Origine", "saisi_par": "Saisi par", "saisi_le": "Saisi le",
                            "releve": "Pointée sur le relevé du"}
                   ).to_excel(w, sheet_name="Opérations", index=False)
        for ex in sorted(ops["exercice"].unique()):
            if ex:
                synthese_mensuelle(ops, ex, mois_debut).to_excel(
                    w, sheet_name=f"Synthèse {libelle_exercice(ex, mois_debut)}"[:31], index=False)
        for t, df in tables.items():
            if t != "operations":
                df.to_excel(w, sheet_name=t.replace("_", " ").capitalize()[:31], index=False)
        mettre_en_forme(w)
    return buf.getvalue()


def mettre_en_forme(writer) -> None:
    from openpyxl.styles import Font, PatternFill
    for ws in writer.book.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1A1A1A")
        ws.freeze_panes = "A2"
        for col in ws.columns:
            largeur = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, largeur + 2), 60)
            entete = str(col[0].value).lower()
            if entete in ("montant", "recettes", "dépenses", "résultat", "total", "prévu", "réalisé") or \
                    entete.rstrip(".").capitalize() in [m.rstrip(".") for m in MOIS_COURTS]:
                for c in col[1:]:
                    c.number_format = '#,##0.00 "€"'
